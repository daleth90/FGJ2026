"""檢查 FGJ 2026 表格格式。

用法：python Doc/check_table.py [xlsx 路徑]
預設檢查 Doc/FGJ 2026表格_v2.xlsx。有錯誤時 exit code = 1。
"""
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ATTRS = ["mor", "hmd", "spd", "tgh"]
NON_NEGATIVE = {"hmd", "spd", "tgh"}
OFFSET_ITEM = re.compile(r"^(mor|hmd|spd|tgh):[+-]\d+$")
REQUIRE_ITEM = re.compile(r"^(mor|hmd|spd|tgh)(>=|<=)(-?\d+)$")
# 第幾階段（event_id 千位數）→ 哪些選項必須減少濕度
TIER_MUST_REDUCE_HMD = {1: "c", 2: "bc", 3: "abc"}

errors, warnings = [], []


def err(where, msg):
    errors.append(f"{where}: {msg}")


def warn(where, msg):
    warnings.append(f"{where}: {msg}")


def read_sheet(ws):
    """以第 2 列的 key 當欄位名，回傳 (列號, dict) 清單。"""
    keys = [c.value for c in ws[2]]
    rows = []
    for r in range(4, ws.max_row + 1):
        values = [ws.cell(r, c + 1).value for c in range(len(keys))]
        if all(v is None for v in values):
            continue
        rows.append((r, {k: v for k, v in zip(keys, values) if k}))
    return rows


def as_int(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return None


def parse_offset(where, s):
    if s is None:
        return {}
    if not isinstance(s, str) or s != s.strip() or " " in s:
        err(where, f"offset 不可有空白：{s!r}")
        return {}
    result = {}
    for part in s.split(";"):
        if not OFFSET_ITEM.match(part):
            err(where, f"offset 格式錯誤 {part!r}（應為 hmd:-20，正號也要寫）")
            continue
        key, value = part.split(":")
        if key in result:
            err(where, f"offset 屬性重複：{key}")
        result[key] = int(value)
    order = [k for k in ATTRS if k in result]
    if list(result) != order:
        warn(where, f"offset 建議依 mor;hmd;spd;tgh 排序：{s}")
    return result


def parse_require(where, s):
    if s is None:
        return
    if not isinstance(s, str) or " " in s:
        err(where, f"require 不可有空白：{s!r}")
        return
    seen = Counter()
    for part in s.split(";"):
        m = REQUIRE_ITEM.match(part)
        if not m:
            err(where, f"require 格式錯誤 {part!r}（應為 spd>=30 或 mor<=-20）")
            continue
        key, op, value = m.group(1), m.group(2), int(m.group(3))
        seen[(key, op)] += 1
        if key in NON_NEGATIVE and value < 0:
            err(where, f"{key} 不會是負數，需求 {part} 沒有意義")
    for (key, op), n in seen.items():
        if n > 1:
            err(where, f"require 重複條件：{key}{op}")


def check_titles(rows):
    ids = set()
    combos = Counter()
    for r, row in rows:
        where = f"title_list 第{r}列"
        tid = as_int(row.get("title_id"))
        if tid is None:
            err(where, f"title_id 不是整數：{row.get('title_id')!r}")
            continue
        if tid in ids:
            err(where, f"title_id 重複：{tid}")
        ids.add(tid)
        reach = [row.get(f"reach_{a}") for a in ATTRS]
        if all(v is None for v in reach):
            continue  # 特殊結局
        if any(as_int(v) not in (0, 1) for v in reach):
            err(where, f"reach_* 只能填 0 或 1：{reach}")
            continue
        combos[tuple(int(v) for v in reach)] += 1
    for combo, n in combos.items():
        if n > 1:
            err("title_list", f"達標組合重複 {n} 次：{combo}")
    if len(combos) != 2 ** len(ATTRS):
        err("title_list", f"達標組合應有 16 種，目前 {len(combos)} 種")
    return ids


def check_events(rows, title_ids):
    ids = set()
    empty = Counter()
    for r, row in rows:
        where = f"event_list 第{r}列"
        eid = as_int(row.get("event_id"))
        if eid is None:
            err(where, f"event_id 不是整數：{row.get('event_id')!r}")
            continue
        where = f"event {eid}"
        if eid in ids:
            err(where, "event_id 重複")
        ids.add(eid)
        if as_int(row.get("unlock_event_num")) is None:
            err(where, "unlock_event_num 不是整數")
        if not row.get("event_desc"):
            empty["event_desc"] += 1

        tier = eid // 1000
        for k in "abc":
            w = f"{where} 選項{k.upper()}"
            if not row.get(f"choice_{k}_desc"):
                empty[f"choice_{k}_desc"] += 1
            parse_require(w, row.get(f"choice_{k}_require"))
            offset = parse_offset(w, row.get(f"choice_{k}_offset"))

            end = row.get(f"choice_{k}_end")
            if end is not None:
                end_id = as_int(end)
                if end_id is None:
                    err(w, f"end 不是整數：{end!r}")
                elif end_id not in title_ids:
                    err(w, f"end {end_id} 在 title_list 找不到")

            if row.get(f"choice_{k}_desc") and end is None:
                if k in TIER_MUST_REDUCE_HMD.get(tier, "") and offset.get("hmd", 0) >= 0:
                    err(w, f"第{tier}階段此選項必須減少濕度（offset 需含 hmd:-xx）")
                if not offset:
                    warn(w, "有文案但沒有 offset")
    return empty


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("FGJ 2026表格_v2.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    title_ids = check_titles(read_sheet(wb["title_list"]))
    empty = check_events(read_sheet(wb["event_list"]), title_ids)

    print(f"檢查：{path.name}")
    for e in errors:
        print(f"  [錯誤] {e}")
    for w in warnings:
        print(f"  [提醒] {w}")
    if empty:
        print("  尚未填寫：" + "、".join(f"{k} {n} 格" for k, n in empty.items()))
    print(f"結果：{len(errors)} 個錯誤、{len(warnings)} 個提醒")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
